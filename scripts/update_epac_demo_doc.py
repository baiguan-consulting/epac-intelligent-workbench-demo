from copy import deepcopy
from docx import Document


DOCUMENT_PATH = "C:/Users/12359/Desktop/epac演示.docx"


def insert_after(paragraph, text):
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph.style = paragraph.style
    if paragraph.runs:
        reference_run = paragraph.runs[0]
        new_run = new_paragraph.add_run(text)
        new_run.bold = reference_run.bold
        new_run.italic = reference_run.italic
        new_run.font.name = reference_run.font.name
        new_run.font.size = reference_run.font.size
    else:
        new_paragraph.add_run(text)
    paragraph._p.addnext(new_paragraph._p)
    return new_paragraph


document = Document(DOCUMENT_PATH)
paragraphs = document.paragraphs

admin_paragraph = next(
    paragraph
    for paragraph in paragraphs
    if paragraph.text.startswith("行政工作台则把请假、合同用印、招聘、入职和离职流程集中管理。")
)

new_paragraphs = [
    "行政端新增了“行政文件数据上传”和“会议记录”两个入口。行政人员可以把制度、政府公文、合同模板、历史资料等文件统一打包上传；系统会模拟识别文件类型并自动分类归档。",
    "演示时可以点击“会议记录”，再点击“开始模拟录音”。系统会展示 AI 同时识别越南语和中文发言，并把会议内容整理成文字版会议纪要。",
    "AI 不只转写录音，还会自动提炼会议结论，并生成每一项待办的责任人和截止时间。确认后，这份纪要会归档至行政资料库，同时生成待办，后续可以继续追踪。",
    "例如可以直接问智能助手：“上次安全培训会议的待办有哪些？”或“越南员工请假制度是什么？”。系统会从已归档的行政文件和历史会议中给出答案。",
    "这个功能的价值是把原来分散在录音、聊天记录和个人文件夹里的行政资料沉淀下来。即使行政人员离职或新员工入职，也可以快速查到历史制度、会议决定和未完成事项。",
]

anchor = admin_paragraph
for text in new_paragraphs:
    anchor = insert_after(anchor, text)

for paragraph in document.paragraphs:
    if paragraph.text == "行政员工只看请假、合同用印、招聘、入职和离职等行政流程。":
        paragraph.text = "行政员工只看请假、合同用印、招聘、入职和离职，以及行政文件和会议待办等职责范围内的提醒。"
    elif paragraph.text == "销售员用越南语完成客户录入；系统自动生成中文管理记录；销售负责人看到团队推进情况；老板用中文直接提问经营问题；采购和行政在自己的角色范围内完成协同和流程处理。":
        paragraph.text = "销售员用越南语完成客户录入；系统自动生成中文管理记录；销售负责人看到团队推进情况；老板用中文直接提问经营问题；采购在价格和交期上协同，行政则将流程、文件和会议纪要沉淀为可追溯的资料库。"

document.save(DOCUMENT_PATH)
